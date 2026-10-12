"""Build uat/uat_test_cases.xlsx and the UAT tables in docs/07_uat_plan.md from one source of truth,
and check that every test case matches the traceability matrix in the BRD."""
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from docblocks import md_table, replace_block

ROOT = Path(__file__).resolve().parents[1]

# id, title, requirement, story, preconditions, steps, expected, cycle 1, cycle 2, defect
CASES = [
    ("TC-01", "Valid customer number auto-fills customer data", "FR-01", "US-01", "Customer 100245 exists in SAP QA",
     "Create complaint; enter customer no. 100245", "Name, country, account manager, type filled and read-only", "Pass", "-", ""),
    ("TC-02", "Invalid customer number rejected", "FR-01", "US-01", "-",
     "Enter customer no. 999999; click Save", "'Customer not found in SAP'; record not saved", "Pass", "-", ""),
    ("TC-03", "Valid material code auto-fills material data", "FR-02", "US-02", "AS-4410-22 active in SAP QA",
     "Enter material code AS-4410-22", "Description, product family 'Air springs', plant filled", "Pass", "-", ""),
    ("TC-04", "Obsolete material triggers warning", "FR-02", "US-02", "Material BC-2001-05 status obsolete",
     "Enter material code BC-2001-05", "Warning shown; user must confirm to continue", "Fail", "Pass", "DEF-01"),
    ("TC-05", "Save blocked when batch number missing", "FR-03", "US-03", "-",
     "Complete all mandatory fields except batch; click Save", "Not saved; batch field highlighted with message", "Pass", "-", ""),
    ("TC-06", "Save succeeds with all mandatory fields", "FR-03", "US-03", "-",
     "Complete all mandatory fields; click Save", "Record saved; complaint ID generated", "Pass", "-", ""),
    ("TC-07", "Defect picklist filtered by product family", "FR-04", "US-04", "Product family 'Brake chambers'",
     "Open defect code list", "Only brake-chamber defect codes listed", "Fail", "Pass", "DEF-05"),
    ("TC-08", "Free text not accepted as defect code", "FR-04", "US-04", "-",
     "Type a new value into defect code field", "Value rejected; free text only in description", "Pass", "-", ""),
    ("TC-09", "Duplicate customer reference warning", "FR-05", "US-05", "CMP-01200 exists with ref QN-55871",
     "Register new complaint, same customer and reference", "'Possible duplicate of CMP-01200'; Link / Continue", "Pass", "-", ""),
    ("TC-10", "Auto-assignment for mapped product family", "FR-06", "US-06", "Air springs mapped to Assembly Line 2",
     "Log first response on an air-spring complaint", "Assigned to Assembly Line 2 owner and engineer", "Pass", "-", ""),
    ("TC-11", "Unmapped product family goes to team lead queue", "FR-06", "US-06", "Test family with no owner",
     "Log first response", "Case appears in team lead queue", "Pass", "-", ""),
    ("TC-12", "Reminder at 1 day without first response", "FR-07", "US-07", "Complaint received Mon 09:00",
     "Advance test clock to Tue 09:00", "Reminder e-mail to assigned engineer", "Pass", "-", ""),
    ("TC-13", "Escalation at 1.5 days without first response", "FR-07", "US-07", "Engineer reassigned after creation",
     "Advance test clock to Tue 21:00", "Escalation e-mail to the current engineer's team lead", "Fail", "Pass", "DEF-02"),
    ("TC-14", "Cannot enter D5 without verified root cause", "FR-08", "US-08", "Case in D4, root cause not verified",
     "Move case to D5", "Blocked with message", "Pass", "-", ""),
    ("TC-15", "Closure requires D8 fields and customer confirmation", "FR-08", "US-08", "Case in D8",
     "Close case without customer confirmation", "Closure blocked until confirmation recorded", "Pass", "-", ""),
    ("TC-16", "Overdue corrective action notification", "FR-09", "US-09", "Action due yesterday, open",
     "Run daily alert job", "Action owner and department owner notified", "Pass", "-", ""),
    ("TC-17", "Credit note EUR 1,500 routed to team lead only", "FR-10", "US-10", "-",
     "Request credit note EUR 1,500", "Single approval step: team lead", "Pass", "-", ""),
    ("TC-18", "Credit note EUR 4,800 routed to QM then Finance", "FR-10", "US-10", "-",
     "Request credit note EUR 4,800", "Two steps: Quality Manager, then Finance Controller", "Pass", "-", ""),
    ("TC-19", "Rejected credit note returns with mandatory comment", "FR-10", "US-10", "Request pending with team lead",
     "Reject without comment; then with comment", "Blocked without comment; returned to requester with comment", "Pass", "-", ""),
    ("TC-20", "Approval audit trail recorded and locked", "FR-11", "US-11", "Approved credit note",
     "Open audit trail; try to edit as admin", "Approver, timestamp, decision, comment shown; edit not possible", "Pass", "-", ""),
    ("TC-21", "Approved credit note in SAP hand-off file", "FR-12", "US-12", "Two approvals today",
     "Generate daily hand-off file; import in SAP QA", "Two rows; dot decimal; import succeeds", "Fail", "Pass", "DEF-03"),
    ("TC-22", "Weekly KPI extract reconciles with case list", "FR-13", "US-13", "Test week with 60 cases, 2 reopened",
     "Run weekly extract; compare with case list", "All KPI figures match case list", "Fail", "Pass", "DEF-04"),
    ("TC-23", "Attach three photos up to 10 MB", "FR-14", "US-14", "-",
     "Attach 3 photos (4, 8, 10 MB)", "All attached and visible to assigned engineer", "Pass", "-", ""),
    ("TC-24", "Engineer role cannot approve credit notes", "NFR-02", "US-10", "User with engineer role",
     "Open pending credit note", "Approve/Reject buttons not available", "Pass", "-", ""),
    ("TC-25", "SAP look-up response time", "NFR-01", "US-01, US-02", "Business hours, 20 look-ups",
     "Time customer and material look-ups", "95% of look-ups within 2 seconds", "Pass", "-", ""),
]
COLUMNS = ["Test ID", "Title", "Requirement", "User story", "Preconditions", "Steps", "Expected result",
           "Cycle 1", "Cycle 2 (retest)", "Defect ID"]

DEFECTS = [
    ("DEF-01", "TC-04", "Obsolete material accepted without warning (status field not mapped)", "Major", "Fixed - retest passed"),
    ("DEF-02", "TC-13", "Escalation sent to original engineer's lead after reassignment", "Major", "Fixed - retest passed"),
    ("DEF-03", "TC-21", "Hand-off file used comma decimal separator; SAP import rejected rows", "Critical", "Fixed - retest passed"),
    ("DEF-04", "TC-22", "Reopened cases counted twice in weekly received volume", "Major", "Fixed - retest passed"),
    ("DEF-05", "TC-07", "Two retired defect codes still shown for Brake chambers", "Minor", "Fixed - retest passed"),
]
DEFECT_COLUMNS = ["Defect ID", "Test ID", "Description", "Severity", "Status"]


def check_traceability(cases):
    """Every (test case, requirement) pair must appear in the BRD traceability matrix and vice versa."""
    brd = (ROOT / "docs/05_requirements_brd.md").read_text()
    matrix = brd.split("## 5. Requirements traceability matrix")[1]
    in_brd = set()
    for line in matrix.splitlines():
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) == 4 and re.match(r"N?FR-\d+", cells[1]):
            in_brd |= {(tc, cells[1]) for tc in re.findall(r"TC-\d+", cells[3])}
    in_cases = {(c[0], c[2]) for c in cases}
    assert in_brd == in_cases, f"traceability mismatch: {in_brd ^ in_cases}"


def write_xlsx(cases, defects, path):
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    status_fill = {"Pass": "D9EAD3", "Fail": "F4CCCC"}
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, df, widths in [("Test cases", cases, [9, 44, 12, 12, 30, 40, 46, 9, 14, 10]),
                                 ("Defect log", defects, [10, 9, 60, 10, 22])]:
            df.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            for cell in ws[1]:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), header_fill
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.value in status_fill:
                        cell.fill = PatternFill("solid", fgColor=status_fill[cell.value])
        summary = xl.book.create_sheet("Summary", 0)
        final = cases["Cycle 2 (retest)"].where(cases["Cycle 2 (retest)"] != "-", cases["Cycle 1"])
        rows = [("UAT summary - Complaint handling redesign", ""), ("Test cases executed", len(cases)),
                ("Passed in cycle 1", int((cases["Cycle 1"] == "Pass").sum())),
                ("Failed in cycle 1", int((cases["Cycle 1"] == "Fail").sum())),
                ("Passed after retest", int((final == "Pass").sum())), ("Defects raised", len(defects)),
                ("Open critical / major defects", 0), ("Exit criteria met", "Yes")]
        for r in rows:
            summary.append(r)
        summary["A1"].font = Font(bold=True, size=13)
        summary.column_dimensions["A"].width = 36


def main():
    cases, defects = pd.DataFrame(CASES, columns=COLUMNS), pd.DataFrame(DEFECTS, columns=DEFECT_COLUMNS)
    check_traceability(CASES)
    assert set(cases["Defect ID"]) - {""} == set(defects["Defect ID"])
    assert (cases.loc[cases["Cycle 1"] == "Fail", "Cycle 2 (retest)"] == "Pass").all()

    out = ROOT / "uat"
    out.mkdir(exist_ok=True)
    write_xlsx(cases, defects, out / "uat_test_cases.xlsx")

    doc = ROOT / "docs/07_uat_plan.md"
    replace_block(doc, "TESTCASES", md_table(cases[["Test ID", "Title", "Requirement", "User story", "Cycle 1",
                                                    "Cycle 2 (retest)", "Defect ID"]]))
    replace_block(doc, "DEFECTS", md_table(defects))
    fails = (cases["Cycle 1"] == "Fail").sum()
    replace_block(doc, "UATSUMMARY", f"**{len(cases)} test cases** executed over two cycles: {len(cases) - fails} passed "
                  f"first time, {fails} failed and passed on retest. **{len(defects)} defects** raised "
                  f"({(defects['Severity'] == 'Critical').sum()} critical, {(defects['Severity'] == 'Major').sum()} major, "
                  f"{(defects['Severity'] == 'Minor').sum()} minor), all fixed and retested before sign-off.")
    print(f"UAT: {len(cases)} cases, {fails} failed in cycle 1, {len(defects)} defects -> uat/uat_test_cases.xlsx")


if __name__ == "__main__":
    main()
